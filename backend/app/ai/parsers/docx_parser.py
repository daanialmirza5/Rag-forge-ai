import io

import docx

from app.ai.parsers.base import DocumentParser, ParsedDocument, ParsedPage


class DocxParser(DocumentParser):
    """`.docx` has no reliable programmatic page-break concept (pagination is
    a rendering-time computation, not stored data), so the whole document is
    treated as a single logical page for citation purposes."""

    @property
    def supported_mime_types(self) -> set[str]:
        return {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}

    def parse(self, file_bytes: bytes) -> ParsedDocument:
        document = docx.Document(io.BytesIO(file_bytes))

        parts: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))

        return ParsedDocument(pages=[ParsedPage(page_number=1, text="\n".join(parts))])
