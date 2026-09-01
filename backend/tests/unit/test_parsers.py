import io

import docx
import pypdf

from app.ai.parsers.docx_parser import DocxParser
from app.ai.parsers.factory import get_parser_for_mime_type, is_supported_mime_type
from app.ai.parsers.html_parser import HTMLParser
from app.ai.parsers.pdf_parser import PDFParser
from app.ai.parsers.text_parser import MarkdownParser, PlainTextParser


def test_plain_text_parser_roundtrip():
    parser = PlainTextParser()
    doc = parser.parse(b"Hello, world!\nSecond line.")
    assert len(doc.pages) == 1
    assert doc.pages[0].page_number == 1
    assert "Hello, world!" in doc.full_text
    assert not doc.is_empty


def test_markdown_parser_strips_syntax():
    parser = MarkdownParser()
    doc = parser.parse(b"# Heading\n\nSome **bold** text and a [link](http://example.com).")
    assert "Heading" in doc.full_text
    assert "Some" in doc.full_text
    assert "**" not in doc.full_text
    assert "](http" not in doc.full_text


def test_html_parser_strips_tags_and_scripts():
    parser = HTMLParser()
    html = b"""
    <html><body>
      <script>alert('x')</script>
      <h1>Title</h1>
      <p>Body paragraph.</p>
    </body></html>
    """
    doc = parser.parse(html)
    assert "Title" in doc.full_text
    assert "Body paragraph." in doc.full_text
    assert "alert" not in doc.full_text


def test_docx_parser_extracts_paragraphs_and_tables():
    document = docx.Document()
    document.add_paragraph("Intro paragraph.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Value"
    buf = io.BytesIO()
    document.save(buf)

    parser = DocxParser()
    parsed = parser.parse(buf.getvalue())
    assert "Intro paragraph." in parsed.full_text
    assert "Name" in parsed.full_text and "Value" in parsed.full_text


def test_pdf_parser_handles_blank_page_without_crashing():
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)

    parser = PDFParser()
    parsed = parser.parse(buf.getvalue())
    # No embedded text and no OCR binaries available in this environment —
    # parsing must degrade gracefully (empty page text), never raise.
    assert len(parsed.pages) == 1
    assert parsed.pages[0].page_number == 1


def test_factory_resolves_parser_by_mime_type():
    assert is_supported_mime_type("text/plain")
    assert is_supported_mime_type("application/pdf")
    assert not is_supported_mime_type("application/x-nonexistent")
    assert isinstance(get_parser_for_mime_type("text/plain"), PlainTextParser)


def test_factory_raises_for_unsupported_mime_type():
    import pytest

    from app.core.exceptions import ValidationAppError

    with pytest.raises(ValidationAppError):
        get_parser_for_mime_type("application/x-nonexistent")
